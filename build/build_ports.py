# -*- coding: utf-8 -*-
"""Ensambla la presentación «Puertos de Comunicación de PC».

Uso:
    python3 build_ports.py              # PPTX + vista previa (tema oscuro)
    python3 build_ports.py --light      # PPTX + vista previa (tema claro)
    python3 build_ports.py --pptx       # solo PPTX
    python3 build_ports.py --preview    # solo imágenes de vista previa
    python3 build_ports.py --pdf        # PDF de vista previa a partir de las imágenes
"""
import os, sys, time, glob
from PIL import Image

import theme
from theme import SLIDE_W, SLIDE_H, SECTION, ROOT
import layouts
from pptx import Presentation
from pptx.util import Inches

import importlib
import ports_content1, ports_content2, ports_content3

CONTENT_MODULES = (ports_content1, ports_content2, ports_content3)

layouts.DECK_TITLE = "Puertos de Comunicación de PC · Guía Completa y Detallada"

# sección -> (acento, número, nombre)
SECTION.clear()
SECTION.update({
    1: ("cyan",   "01", "Puertos de uso general y datos"),
    2: ("orange", "02", "Entrada de periféricos"),
    3: ("purple", "03", "Puertos de video y pantalla"),
    4: ("green",  "04", "Red y comunicaciones"),
    5: ("pink",   "05", "Puertos de audio"),
    6: ("blue",   "06", "Almacenamiento y expansión"),
    7: ("teal",   "07", "Infrarrojos y repaso final"),
})

OUT_DARK = os.path.join(ROOT, "Puertos de Comunicación de PC - Guía Completa.pptx")
OUT_LIGHT = os.path.join(ROOT, "Puertos de Comunicación de PC - Guía Completa (Fondo Claro).pptx")
PREVIEW_DARK = os.path.join(ROOT, "_preview_puertos")
PREVIEW_LIGHT = os.path.join(ROOT, "_preview_puertos_claro")
PDF_DARK = os.path.join(ROOT, "Puertos de Comunicación de PC - Vista Previa.pdf")
PDF_LIGHT = os.path.join(ROOT, "Puertos de Comunicación de PC - Vista Previa (Fondo Claro).pdf")

DECK_FOOT = "Puertos de Comunicación de PC · Guía Completa y Detallada"


def reload_content():
    for m in CONTENT_MODULES:
        importlib.reload(m)


def build_specs():
    reload_content()
    specs = []
    for m in CONTENT_MODULES:
        specs += m.SLIDES
    for s in specs:
        s.setdefault("foot", DECK_FOOT + " · 2026")
    return specs


def meta_for(spec):
    acc, num, name = SECTION[spec.get("section", 1)]
    return {"accent": acc, "num": num, "name": name}


# --------------------------------------------------------------------- PPTX
def build_pptx(specs, path):
    from canvas import PptxCanvas
    import deck as deckmod
    import anim
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    total = len(specs)
    for i, spec in enumerate(specs, 1):
        slide = prs.slides.add_slide(blank)
        meta = meta_for(spec)
        c = PptxCanvas(slide, SLIDE_W, SLIDE_H, accent=meta["accent"], total=total)
        groups = deckmod.render(c, meta, i, total, spec)
        anim.set_transition(slide, spec.get("transition", "fade"))
        if spec.get("anim", True):
            anim.add_animation(slide, groups)
        anim.set_notes(slide, spec.get("notes", ""))
        if i % 15 == 0:
            print("   ... %d/%d diapositivas" % (i, total))
    core = prs.core_properties
    core.title = "Puertos de Comunicación de PC — Guía Completa y Detallada"
    core.subject = "Hardware · Identificación de puertos y especificaciones"
    core.author = "Guía completa de hardware"
    core.comments = ("Guía completa de 23 puertos de comunicación con diagramas de "
                     "identificación, códigos de colores y especificaciones, en %d diapositivas "
                     "con transiciones y animaciones." % total)
    prs.save(path)
    return total


# ------------------------------------------------------------------ PREVIEW
def build_preview(specs, folder):
    from canvas import PilCanvas
    import deck as deckmod
    os.makedirs(folder, exist_ok=True)
    for f in os.listdir(folder):
        if f.endswith((".png", ".jpg")):
            os.remove(os.path.join(folder, f))
    total = len(specs)
    for i, spec in enumerate(specs, 1):
        meta = meta_for(spec)
        c = PilCanvas(SLIDE_W, SLIDE_H, total=total, accent=meta["accent"])
        deckmod.render(c, meta, i, total, spec)
        img = c.finish()
        img.save(os.path.join(folder, "slide_%03d.jpg" % i), quality=82, optimize=True)
        if i % 15 == 0:
            print("   ... preview %d/%d" % (i, total))
    return total


def make_pdf(folder, path):
    files = sorted(glob.glob(os.path.join(folder, "slide_*.jpg")))
    if not files:
        print("   (sin imágenes para el PDF)")
        return
    imgs = [Image.open(f).convert("RGB") for f in files]
    imgs[0].save(path, save_all=True, append_images=imgs[1:], resolution=130.0)
    print("PDF: %s (%d páginas)" % (path, len(imgs)))


def make_contact_sheet(folder, path, cols=4):
    files = sorted(glob.glob(os.path.join(folder, "slide_*.jpg")))
    if not files:
        return
    ims = [Image.open(f) for f in files]
    w, h = ims[0].size
    tw, th = w // 4, h // 4
    rows = (len(ims) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * tw, rows * th), (12, 18, 32))
    for i, im in enumerate(ims):
        im = im.resize((tw, th), Image.LANCZOS)
        r, c = divmod(i, cols)
        sheet.paste(im, (c * tw, r * th))
    sheet.save(path, quality=80, optimize=True)
    print("Índice general: %s" % path)


def main():
    t0 = time.time()
    raw = sys.argv[1:]
    args = [a for a in raw if a not in ("--light", "--dark")]
    target = "light" if "--light" in raw else "dark"
    theme.set_palette(target)
    theme.set_gen("gen_ports_light" if target == "light" else "gen_ports")
    out = OUT_LIGHT if target == "light" else OUT_DARK
    prev = PREVIEW_LIGHT if target == "light" else PREVIEW_DARK
    pdf = PDF_LIGHT if target == "light" else PDF_DARK

    specs = build_specs()
    total = len(specs)
    print("Paleta: %s | diapositivas: %d" % (target, total))
    do_pptx = ("--pptx" in args) or not args
    do_prev = ("--preview" in args) or not args
    do_pdf = ("--pdf" in args) or (not args)
    if do_pptx:
        build_pptx(specs, out)
        print("PPTX escrito en: %s" % out)
    if do_prev:
        build_preview(specs, prev)
        make_contact_sheet(prev, os.path.join(prev, "indice_general.jpg"))
        print("Vista previa en: %s" % prev)
    if do_pdf:
        if not do_prev and not os.path.isdir(prev):
            build_preview(specs, prev)
        make_pdf(prev, pdf)
    print("Listo en %.1f s" % (time.time() - t0))


if __name__ == "__main__":
    main()
