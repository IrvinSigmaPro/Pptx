# -*- coding: utf-8 -*-
"""Primitivas de dibujo con doble backend: PPTX (python-pptx) y PIL (vista previa).

La capa de diseño (layouts.py) trabaja contra esta interfaz, de modo que la
vista previa en PNG usa exactamente la misma geometría que el PPTX final.
"""
import os, re
from theme import COL, rgbname, PXIN, IN, FONT_H, FONT_B, FONT_M, hexcol, TTF

from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

from PIL import Image, ImageDraw, ImageFont, ImageFilter

# ------------------------------------------------------------------ utilidades
TAGS = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def parse_markup(s):
    """Convierte marcado ligero en una lista de runs (texto, bold, color_key, mono)."""
    runs = []
    for part in TAGS.split(s):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            runs.append((part[2:-2], True, None, False))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            runs.append((part[1:-1], False, "@mono", True))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            runs.append((part[1:-1], True, "@accent", False))
        else:
            runs.append((part, False, None, False))
    return runs


_font_cache = {}


def pil_font(size_pt, bold=False, mono=False):
    key = (round(size_pt, 1), bold, mono)
    if key not in _font_cache:
        if mono:
            path = TTF + ("DejaVuSansMono-Bold.ttf" if bold else "DejaVuSansMono.ttf")
        else:
            path = TTF + ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf")
        _font_cache[key] = ImageFont.truetype(path, int(round(size_pt * PXIN / 72.0)))
    return _font_cache[key]


def auto_mode(path, w, h):
    """contain para diagramas generados o imágenes con forma muy distinta a la caja."""
    try:
        from PIL import Image as _I
        iw, ih = _I.open(path).size
    except Exception:
        return "cover"
    ar_img = iw / float(ih)
    ar_box = w / float(h)
    r = ar_img / ar_box
    if "gen/" in path.replace("\\", "/") or r > 1.35 or r < 0.78:
        return "contain"
    return "cover"


def _run_width(txt, size_pt, bold, mono):
    """Ancho en pulgadas."""
    px = pil_font(size_pt, bold, mono).getlength(txt)
    return px / PXIN


def wrap_runs(runs, width_in, size_pt):
    """Ajuste de línea respetando runs. Devuelve lista de líneas (lista de runs)."""
    # tokenizar en palabras conservando estilo
    words = []
    for txt, bold, ck, mono in runs:
        pieces = re.split(r"(\s+)", txt)
        for p in pieces:
            if p == "":
                continue
            words.append((p, bold, ck, mono, p.isspace()))
    lines, cur, curw = [], [], 0.0
    space = _run_width(" ", size_pt, False, False)
    for w, bold, ck, mono, is_space in words:
        wpx = _run_width(w, size_pt, bold, mono)
        if is_space:
            if cur:
                cur.append((w, bold, ck, mono))
                curw += wpx
            continue
        if cur and curw + wpx > width_in + 0.001:
            while cur and cur[-1][0].isspace():
                curw -= _run_width(cur[-1][0], size_pt, cur[-1][1], cur[-1][3])
                cur.pop()
            lines.append(cur)
            cur, curw = [(w, bold, ck, mono)], wpx
        else:
            cur.append((w, bold, ck, mono))
            curw += wpx
    if cur:
        lines.append(cur)
    return lines or [[]]


def measure(text, width_in, size_pt, bold=False, mono=False, spacing=1.12):
    """(alto en pulgadas, nº de líneas) de un texto dentro de un ancho dado."""
    runs = parse_markup(text) if isinstance(text, str) else text
    lines = wrap_runs([(t, b or bold, c, m or mono) for t, b, c, m in runs], width_in, size_pt)
    line_h = size_pt * 1.21 * spacing / 72.0
    return len(lines) * line_h, len(lines)


def autosize(text, width_in, height_in, size_pt, bold=False, mono=False, spacing=1.12, min_size=8.5, floor=0.0):
    """Reduce el tamaño de fuente hasta que el texto quepa en la caja."""
    s = size_pt
    while s > max(min_size, floor):
        h, _ = measure(text, width_in, s, bold, mono, spacing)
        if h <= height_in + 0.002:
            break
        s -= 0.5
    return s


# ==========================================================================
#  Backend PIL (vista previa)
# ==========================================================================
class PilCanvas:
    """Dibuja en un lienzo RGB de 1280x720 px (13.333x7.5 in @96dpi)."""

    def __init__(self, width_in, height_in, total=1, accent="cyan"):
        self.accent = accent
        self.W, self.H = int(width_in * PXIN), int(height_in * PXIN)
        self.width_in, self.height_in = width_in, height_in
        self.base = Image.new("RGB", (self.W, self.H), hexcol("bg"))
        self.layer = Image.new("RGBA", (self.W, self.H), (0, 0, 0, 0))
        self.d = ImageDraw.Draw(self.layer)
        self.elements = []
        self.total = total
        self.hidden = set()       # ids que no se dibujan (usado por el generador de GIF)
        self._id = 1              # la primera forma del PPTX recibe el id 2

    # -- helpers
    def px(self, v):
        return v * PXIN

    def _rgba(self, c, alpha=None):
        v = hexcol(c, None)
        if alpha is None:
            return v + (255,)
        return v + (int(alpha * 255),)

    def _accent(self, key, fallback="cyan"):
        return self.accent if key == "@accent" else (key or fallback)

    def visible(self):
        """Avanza el contador (igual que los shape_id del PPTX) y dice si se dibuja."""
        self._id += 1
        return self._id not in self.hidden

    def register(self, x, y, w, h, kind, text=""):
        el = {"id": self._id, "x": x, "y": y, "w": w, "h": h, "kind": kind, "text": text}
        self.elements.append(el)
        return el

    # -- primitivas
    def blob(self, cx, cy, r, color, alpha=0.20):
        if not self.visible():
            return None
        img = Image.new("RGBA", (int(r * 2 * PXIN), int(r * 2 * PXIN)), (0, 0, 0, 0))
        dr = ImageDraw.Draw(img)
        dr.ellipse([0, 0, img.width - 1, img.height - 1], fill=self._rgba(color, alpha))
        img = img.filter(ImageFilter.GaussianBlur(radius=r * PXIN * 0.45))
        self.layer.alpha_composite(img, (int(cx * PXIN - img.width / 2), int(cy * PXIN - img.height / 2)))

    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.0, alpha=None, kind="rect", text=""):
        if not self.visible():
            return None
        xy = [self.px(x), self.px(y), self.px(x + w), self.px(y + h)]
        rad = self.px(radius)
        if fill:
            self.d.rounded_rectangle(xy, radius=rad, fill=self._rgba(fill, alpha))
        if line:
            self.d.rounded_rectangle(xy, radius=rad, outline=self._rgba(line), width=max(1, int(self.px(lw / 72.0))))
        return self.register(x, y, w, h, kind, text)

    def circle(self, cx, cy, r, fill=None, line=None, lw=1.0, alpha=None):
        if not self.visible():
            return None
        xy = [self.px(cx - r), self.px(cy - r), self.px(cx + r), self.px(cy + r)]
        if fill:
            self.d.ellipse(xy, fill=self._rgba(fill, alpha))
        if line:
            self.d.ellipse(xy, outline=self._rgba(line), width=max(1, int(self.px(lw / 72.0))))
        return self.register(cx - r, cy - r, 2 * r, 2 * r, "circle")

    def line(self, x1, y1, x2, y2, color="line", lw=1.5, dash=False, alpha=None):
        if not self.visible():
            return None
        if dash:
            n = max(2, int(((x2 - x1) ** 2 + (y2 - y1) ** 2) ** 0.5 / 0.12))
            for i in range(n):
                t0, t1 = i / n, (i + 0.55) / n
                self.d.line([self.px(x1 + (x2 - x1) * t0), self.px(y1 + (y2 - y1) * t0),
                             self.px(x1 + (x2 - x1) * t1), self.px(y1 + (y2 - y1) * t1)],
                            fill=self._rgba(color, alpha), width=max(1, int(self.px(lw / 72.0))))
        else:
            self.d.line([self.px(x1), self.px(y1), self.px(x2), self.px(y2)],
                        fill=self._rgba(color, alpha), width=max(1, int(self.px(lw / 72.0))))
        return None

    def text(self, x, y, w, h, s, size=14, color="txt2", bold=False, font="body",
             align="l", anchor="t", spacing=1.12, alpha=None, shrink=True, kind="text"):
        if not s:
            return None
        if not self.visible():
            return None
        runs = parse_markup(s) if isinstance(s, str) else s
        size = autosize(runs, w, h, size, bold, font == "mono", spacing) if shrink else size
        lines = wrap_runs([(t, b or bold, c, m or (font == "mono")) for t, b, c, m in runs], w, size)
        lh = size * 1.21 * spacing / 72.0
        total = len(lines) * lh
        if anchor == "m":
            cy = y + (h - total) / 2.0
        elif anchor == "b":
            cy = y + h - total
        else:
            cy = y
        for ln in lines:
            wl = sum(_run_width(t, size, b, m) for t, b, c, m in ln)
            if align == "c":
                cx = x + (w - wl) / 2.0
            elif align == "r":
                cx = x + w - wl
            else:
                cx = x
            for t, b, c, m in ln:
                col = self._accent(c, None)
                if c == "@mono":
                    col = "yellow"
                elif c is None:
                    col = color
                f = pil_font(size, b, m)
                self.d.text((self.px(cx), self.px(cy + 0.012)), t, font=f, fill=self._rgba(col, alpha))
                cx += _run_width(t, size, b, m)
            cy += lh
        return self.register(x, y, w, h, kind, s)

    def image(self, path, x, y, w, h, mode="cover", radius=0.0, border=None):
        if not self.visible():
            return None
        im = Image.open(path).convert("RGB")
        if mode == "auto":
            mode = auto_mode(path, w, h)
        tw, th = int(self.px(w)), int(self.px(h))
        if mode == "cover":
            sc = max(tw / im.width, th / im.height)
            im = im.resize((max(1, int(im.width * sc)), max(1, int(im.height * sc))), Image.LANCZOS)
            left = (im.width - tw) // 2
            top = (im.height - th) // 3
            im = im.crop((left, top, left + tw, top + th))
        else:
            sc = min(tw / im.width, th / im.height)
            im2 = Image.new("RGB", (tw, th), hexcol("panel"))
            r = im.resize((max(1, int(im.width * sc)), max(1, int(im.height * sc))), Image.LANCZOS)
            im = im2
            im.paste(r, ((tw - r.width) // 2, (th - r.height) // 2))
        mask = Image.new("L", (tw, th), 255)
        if radius:
            ImageDraw.Draw(mask).rounded_rectangle([0, 0, tw - 1, th - 1], radius=self.px(radius), fill=255)
        self.layer.paste(im.convert("RGBA"), (int(self.px(x)), int(self.px(y))), mask)
        if border:
            self.d.rounded_rectangle([self.px(x), self.px(y), self.px(x + w), self.px(y + h)],
                                     radius=self.px(radius), outline=self._rgba(border), width=2)
        return self.register(x, y, w, h, "image", os.path.basename(path))

    def finish(self):
        out = Image.alpha_composite(self.base.convert("RGBA"), self.layer).convert("RGB")
        return out


# ==========================================================================
#  Backend PPTX
# ==========================================================================
class PptxCanvas:
    def __init__(self, slide, width_in, height_in, accent="cyan", total=1):
        self.accent = accent
        self.slide = slide
        self.width_in, self.height_in = width_in, height_in
        self.elements = []
        self.total = total
        self.shapes = slide.shapes

    def px(self, v):
        return v

    def _rgba(self, c, alpha=None):
        return hexcol(c, None)

    def _accent(self, key, fallback="cyan"):
        return self.accent if key == "@accent" else (key or fallback)

    def register(self, shape, x, y, w, h, kind, text=""):
        el = {"id": shape.shape_id, "shape": shape, "x": x, "y": y, "w": w, "h": h,
              "kind": kind, "text": text}
        self.elements.append(el)
        return el

    # -- primitivas
    def blob(self, cx, cy, r, color, alpha=0.20):
        sh = self.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
        sh.shadow.inherit = False
        sh.line.fill.background()
        _grad_oval(sh, color)
        return self.register(sh, cx - r, cy - r, 2 * r, 2 * r, "blob")

    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, radius=0.0, alpha=None, kind="rect", text=""):
        shp = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sh = self.shapes.add_shape(shp, Inches(x), Inches(y), Inches(w), Inches(h))
        if radius:
            _set_radius(sh, min(0.5, radius / max(w, h) if min(w, h) else 0.08))
        sh.shadow.inherit = False
        if fill:
            _solid(sh, fill, alpha)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgbname(line)
            sh.line.width = Pt(max(0.5, lw))
        else:
            sh.line.fill.background()
        sh.text_frame.text = ""
        return self.register(sh, x, y, w, h, kind, text)

    def circle(self, cx, cy, r, fill=None, line=None, lw=1.0, alpha=None):
        sh = self.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
        sh.shadow.inherit = False
        if fill:
            _solid(sh, fill, alpha)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgbname(line)
            sh.line.width = Pt(max(0.5, lw))
        else:
            sh.line.fill.background()
        return self.register(sh, cx - r, cy - r, 2 * r, 2 * r, "circle")

    def line(self, x1, y1, x2, y2, color="line", lw=1.5, dash=False, alpha=None):
        from pptx.enum.shapes import MSO_CONNECTOR
        cn = self.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cn.line.color.rgb = rgbname(color)
        cn.line.width = Pt(max(0.5, lw))
        if dash:
            _dash(cn)
        return self.register(cn, min(x1, x2), min(y1, y2), abs(x2 - x1), abs(y2 - y1), "line")

    def text(self, x, y, w, h, s, size=14, color="txt2", bold=False, font="body",
             align="l", anchor="t", spacing=1.12, alpha=None, shrink=True, kind="text"):
        if not s:
            return None
        tb = self.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        runs = parse_markup(s) if isinstance(s, str) else s
        if shrink:
            size = autosize(runs, w, h, size, bold, font == "mono", spacing)
        p = tf.paragraphs[0]
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT, "j": PP_ALIGN.JUSTIFY}[align]
        p.line_spacing = spacing
        for t, b, c, m in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(size)
            f.bold = bool(b or bold)
            f.name = FONT_M if (m or font == "mono") else (FONT_H if b or bold else FONT_B)
            col = c if c in (None, "@accent") else ("yellow" if c == "@mono" else c)
            f.color.rgb = rgbname(color if col is None or col == "@accent" else col)
        return self.register(tb, x, y, w, h, kind, s)

    def image(self, path, x, y, w, h, mode="cover", radius=0.0, border=None):
        if mode == "auto":
            mode = auto_mode(path, w, h)
        if mode == "contain":
            from PIL import Image as _I
            iw, ih = _I.open(path).size
            ar = iw / float(ih)
            box_ar = w / float(h)
            if ar > box_ar:
                nw, nh = w, w / ar
            else:
                nh, nw = h, h * ar
            pic = self.shapes.add_picture(path, Inches(x + (w - nw) / 2), Inches(y + (h - nh) / 2), Inches(nw), Inches(nh))
        else:
            pic = self.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
            _crop_to_fill(pic, path, w, h)
        if radius:
            _round_pic(pic, min(0.5, radius / max(w, h)))
        if border:
            pic.line.color.rgb = rgbname(border)
            pic.line.width = Pt(1.0)
        return self.register(pic, x, y, w, h, "image", os.path.basename(path))

    def finish(self):
        return None


# ---------------------------------------------------------- helpers de bajo nivel
def _solid(shape, color, alpha=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgbname(color)
    if alpha is not None:
        from pptx.oxml.ns import qn
        try:
            srgb = shape.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
            if srgb is not None:
                srgb.append(srgb.makeelement(qn('a:alpha'), {'val': str(int(alpha * 100000))}))
        except Exception:
            pass


def _grad_oval(shape, color):
    """Resplandor decorativo: óvalo plano con transparencia (compatible con PPT 2007+)."""
    _solid(shape, color, 0.16)


def _set_radius(shape, frac):
    from pptx.oxml.ns import qn
    prst = shape._element.spPr.find(qn('a:prstGeom'))
    if prst is not None:
        av = prst.find(qn('a:avLst'))
        if av is None:
            av = prst.makeelement(qn('a:avLst'), {})
            prst.append(av)
        for gd in list(av):
            av.remove(gd)
        gd = av.makeelement(qn('a:gd'), {'name': 'adj', 'fmla': 'val %d' % int(frac * 100000)})
        av.append(gd)


def _dash(cn):
    from pptx.oxml.ns import qn
    ln = cn.line._get_or_add_ln()
    d = ln.makeelement(qn('a:prstDash'), {'val': 'dash'})
    ln.append(d)


def _crop_to_fill(pic, path, w_in, h_in):
    """Ajusta el recorte para simular object-fit: cover."""
    import copy
    from PIL import Image as _I
    iw, ih = _I.open(path).size
    target = w_in / float(h_in)
    have = iw / float(ih)
    if abs(target - have) < 0.01:
        return
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0.0
    if have > target:            # imagen más ancha -> recortar lados
        frac = (1.0 - target / have) / 2.0
        pic.crop_left = frac
        pic.crop_right = frac
    else:                        # imagen más alta -> recortar arriba/abajo
        frac = (1.0 - have / target) / 2.0
        pic.crop_top = frac
        pic.crop_bottom = frac


def _round_pic(pic, frac):
    from pptx.oxml.ns import qn
    spPr = pic._element.spPr
    prst = spPr.find(qn('a:prstGeom'))
    if prst is None:
        prst = spPr.makeelement(qn('a:prstGeom'), {'prst': 'roundRect'})
        spPr.append(prst)
    else:
        prst.set('prst', 'roundRect')
    av = prst.find(qn('a:avLst'))
    if av is None:
        av = prst.makeelement(qn('a:avLst'), {})
        prst.append(av)
    for gd in list(av):
        av.remove(gd)
    av.append(av.makeelement(qn('a:gd'), {'name': 'adj', 'fmla': 'val %d' % int(frac * 100000)}))
