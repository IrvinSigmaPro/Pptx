# -*- coding: utf-8 -*-
"""Ensambla la presentación completa: PPTX con animaciones + vista previa en imágenes.

Uso:
    python3 build_all.py            # construye el PPTX y la vista previa
    python3 build_all.py --pptx     # solo el PPTX
    python3 build_all.py --preview  # solo las imágenes de vista previa
"""
import os, sys, time
from PIL import Image
import theme
from theme import SLIDE_W, SLIDE_H, SECTION, PREVIEW, ROOT
from pptx import Presentation
from pptx.util import Inches

import importlib
import content1, content2, content3, content4, content5, content6, content7

CONTENT_MODULES = (content1, content2, content3, content4, content5, content6, content7)


def reload_content():
    """Reconstruye las listas de diapositivas con la paleta/carpeta de diagramas activa."""
    for m in CONTENT_MODULES:
        importlib.reload(m)
from canvas import PptxCanvas, PilCanvas
import deck as deckmod
import anim

OUT_DARK = os.path.join(ROOT, "Compuertas Lógicas y Chips Serie 74XX - Versión Mejorada.pptx")
OUT_LIGHT = os.path.join(ROOT, "Compuertas Lógicas y Chips Serie 74XX - Versión Mejorada (Fondo Claro).pptx")
PREVIEW_DARK = os.path.join(ROOT, "_preview")
PREVIEW_LIGHT = os.path.join(ROOT, "_preview_claro")


def build_specs():
    reload_content()
    specs = []
    specs += content1.SLIDES
    specs += content2.SLIDES
    specs += content3.SLIDES + content7.AFTER_3
    specs += content4.SLIDES + content7.AFTER_4
    specs += content5.SLIDES + content7.AFTER_5
    specs += content6.SLIDES
    return specs


def meta_for(spec):
    acc, num, name = SECTION[spec.get("section", 1)]
    return {"accent": acc, "num": num, "name": name}


# --------------------------------------------------------------------- PPTX
def build_pptx(specs, path):
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    total = len(specs)
    for i, spec in enumerate(specs, 1):
        slide = prs.slides.add_slide(blank)
        meta = meta_for(spec)
        c = PptxCanvas(slide, SLIDE_W, SLIDE_H, accent=meta["accent"], total=total)
        groups = deckmod.render(c, meta, i, total, spec)
        anim.set_transition(slide, spec.get("transition", "fade"))
        if spec.get("anim", True):
            anim.add_animation(slide, groups)
        anim.set_notes(slide, spec.get("notes", ""))
        if i % 20 == 0:
            print("   ... %d/%d diapositivas" % (i, total))
    from pptx.oxml.ns import qn
    core = prs.core_properties
    core.title = "Compuertas Lógicas y Chips Serie 74XX"
    core.subject = "Electrónica Digital · Serie 74XX"
    core.author = "Presentación ampliada"
    core.comments = ("Versión ampliada a %d diapositivas con transiciones y animaciones de "
                     "entrada en cada lámina." % total)
    _finalize_ooxml(prs)
    prs.save(path)
    return total




def _finalize_ooxml(prs):
    """Evita el diálogo de reparación de PowerPoint al abrir el PPTX.

    python-pptx deja dos incoherencias que PowerPoint (y Keynote) rechazan:
    - p:sldSz sigue diciendo type="screen4x3" aunque el lienzo es 16:9, y cx
      no coincide con el valor canónico 12192000.
    - al crear notas del expositor agrega la relación notesMaster, pero no
      escribe p:notesMasterIdLst en presentation.xml (ECMA-376).
    """
    from lxml import etree
    from pptx.oxml.ns import qn
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    sldSz = prs.element.find(qn("p:sldSz"))
    if sldSz is not None:
        sldSz.set("cx", "12192000")
        sldSz.set("cy", "6858000")
        sldSz.set("type", "screen16x9")

    rid = None
    for rel in prs.part.rels.values():
        if rel.reltype == RT.NOTES_MASTER:
            rid = rel.rId
            break
    if not rid or prs.element.find(qn("p:notesMasterIdLst")) is not None:
        return
    lst = etree.Element(qn("p:notesMasterIdLst"))
    item = etree.SubElement(lst, qn("p:notesMasterId"))
    item.set(qn("r:id"), rid)
    master = prs.element.find(qn("p:sldMasterIdLst"))
    if master is not None:
        master.addnext(lst)
    else:
        prs.element.insert(0, lst)


# ------------------------------------------------------------------ PREVIEW
def build_preview(specs, folder=None):
    folder = folder or PREVIEW
    os.makedirs(folder, exist_ok=True)
    for f in os.listdir(folder):
        if f.endswith((".png", ".jpg")):
            os.remove(os.path.join(folder, f))
    total = len(specs)
    for i, spec in enumerate(specs, 1):
        meta = meta_for(spec)
        c = PilCanvas(SLIDE_W, SLIDE_H, total=total, accent=meta["accent"])
        deckmod.render(c, meta, i, total, spec)
        img = c.finish()
        img.save(os.path.join(folder, "slide_%03d.jpg" % i), quality=82, optimize=True)
        if i % 20 == 0:
            print("   ... preview %d/%d" % (i, total))
    return total


def make_gifs(specs, folder, which=(13, 84)):
    """Genera GIF animados que reproducen las entradas por clic de una diapositiva."""
    gifs = []
    total = len(specs)
    for idx in which:
        spec = specs[idx - 1]
        meta = meta_for(spec)
        # 1) render completo para conocer los grupos de animación
        c = PilCanvas(SLIDE_W, SLIDE_H, total=total, accent=meta["accent"])
        groups = deckmod.render(c, meta, idx, total, spec)
        all_ids, steps = [], []
        for g in groups:
            steps.append(list(g["ids"]))
            all_ids += list(g["ids"])
        frames = []
        for k in range(-1, len(steps)):
            shown = set()
            for s_ in steps[:k + 1]:
                shown |= set(s_)
            hidden = set(all_ids) - shown
            c2 = PilCanvas(SLIDE_W, SLIDE_H, total=total, accent=meta["accent"])
            c2.hidden = hidden
            deckmod.render(c2, meta, idx, total, spec)
            frames.append(c2.finish().quantize(colors=200, method=Image.MEDIANCUT))
        dur = [260] + [760] * (len(frames) - 1)
        path = os.path.join(folder, "animacion_diapositiva_%03d.gif" % idx)
        frames[0].save(path, save_all=True, append_images=frames[1:], duration=dur,
                       loop=0, optimize=True, disposal=2)
        gifs.append(path)
        print("   GIF: %s (%d pasos)" % (os.path.basename(path), len(steps)))
    return gifs


def main():
    t0 = time.time()
    args = [a for a in sys.argv[1:] if a not in ("--light", "--dark")]
    raw = sys.argv[1:]
    palette = "light" if "--light" in raw else ("dark" if "--dark" in raw else None)
    if palette is None:
        palette = "dark"
    target = "light" if palette == "light" else "dark"
    theme.set_palette(target)
    theme.set_gen("gen_light" if target == "light" else "gen")
    out = OUT_LIGHT if target == "light" else OUT_DARK
    prev = PREVIEW_LIGHT if target == "light" else PREVIEW_DARK

    specs = build_specs()
    total = len(specs)
    print("Paleta: %s | total de diapositivas: %d" % (target, total))
    do_pptx = ("--pptx" in args) or not args
    do_prev = ("--preview" in args) or not args
    do_gif = ("--gif" in args)
    if do_pptx:
        build_pptx(specs, out)
        print("PPTX escrito en: %s" % out)
    if do_prev:
        build_preview(specs, prev)
        print("Vista previa en: %s" % prev)
    if do_gif:
        os.makedirs(prev, exist_ok=True)
        make_gifs(specs, prev)
    print("Listo en %.1f s" % (time.time() - t0))


if __name__ == "__main__":
    main()
